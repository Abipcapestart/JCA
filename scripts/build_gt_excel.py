"""
Builds one Excel workbook consolidating all three GT-drug runs: per-drug
summary, every comparator/outcome/member-state row, every single Bedrock and
Tavily call with cost, per-stage latency, and a pipeline-steps reference sheet
(which stages carry an LLM prompt vs. which are deterministic/API-only).

Reads scripts/gt_eval_output/{drug}.json (produced by run_gt_eval.py).
Writes scripts/gt_eval_output/GT_Phase1_Results.xlsx.
"""

from __future__ import annotations

import json
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

BASE = os.path.dirname(os.path.abspath(__file__))
OUT_DIR = os.path.join(BASE, "gt_eval_output")
DRUGS = ["Tovorafenib", "Lurbinectedin", "Tarlatamab"]

# The prompt_id prefix does NOT always match the orchestrator stage it runs
# in -- verified against jca_phase1/agents/a14_a18_consolidation.py and
# orchestrator.py directly. harmonize_outcomes() and build_outcomes() (whose
# prompts are named a14.* / one a16.*) both execute inside the orchestrator's
# "A15" block; build_comparator()'s two prompts execute inside "A16".
PROMPT_TO_STAGE = {
    "a01.pi_validation": "A1",
    "a02.input_structuring": "A2",
    "a03.scope_facet_normalise": "A3",
    "a04.indication_lock": "A4",
    "a05.area_adjudication": "A5",
    "a06.query_vocabulary": "A6",
    "a08.extraction": "A8",
    "a10.claim_validation": "A10",
    "a11.comparator_identity": "A11",
    "a12.scope_adjudication": "A12",
    "a14.outcome_harmonization": "A15",
    "a16.outcome_rationale": "A15",
    "a16.comparator_rationale": "A16",
    "a16.indication_synthesis": "A16",
}

PIPELINE_STEPS = [
    ("A1", "P&I Validation", "LLM prompt", "a01.pi_validation",
     "FIX/CHECK gate on the raw population+intervention input before anything else runs."),
    ("A2", "Input Structuring", "LLM prompt", "a02.input_structuring",
     "Maps free/partial input onto the field taxonomy; tags each field confirmed/inferred/not_provided."),
    ("A4", "Indication Lock", "API + LLM prompt", "a04.indication_lock",
     "Resolves the licensed indication via the EMA medicine registry (+ search); LLM only for wording synthesis. Emits the LeakageGuard."),
    ("A3", "Population Scope Boundary", "Deterministic + 1 LLM prompt", "a03.scope_facet_normalise",
     "Turns SME population fields into a typed, checkable predicate. Mostly deterministic; one LLM call normalises free-text facets."),
    ("A5", "Therapeutic Area Resolution", "Deterministic, LLM only if ambiguous", "a05.area_adjudication",
     "Keyword match resolves the area outright in the common case (0 LLM calls observed in all 3 runs); LLM tiebreak only fires on genuine ambiguity."),
    ("sources", "Source Inventory Load", "Deterministic", "-",
     "Loads and validates the curated source workbook (sheet-fuzzy matching, hyperlink reading). No LLM, no network beyond the local file."),
    ("A6", "Query Planning", "Deterministic templates + 1 LLM prompt", "a06.query_vocabulary",
     "Deterministic templates compose the actual search queries (reviewable/diffable). One LLM call builds a vocabulary object that by design cannot contain a comparator/drug name."),
    ("A7", "Source-Routed Retrieval", "API only (no LLM)", "-",
     "Tavily search/extract + ClinicalTrials.gov v2 + PubMed E-utilities + EMA, run in parallel across 27 states x source classes. Pure retrieval; no model calls."),
    ("A8", "Evidence Extraction", "LLM prompt", "a08.extraction",
     "The one shared extraction contract every retrieved document goes through. Structured registry records (trials) bypass the LLM entirely and are typed directly from the API."),
    ("A9", "Grounding", "Deterministic", "-",
     "Lexical check: does the extracted value actually appear in the retrieved text? Catches fabrication; ~0 cost."),
    ("A10", "Claim Validation", "LLM prompt + live re-fetch", "a10.claim_validation",
     "Re-opens the cited source fresh (not the extraction-time text window) and asks whether it supports the claim for this population/intervention. The single biggest source of rejections in every run."),
    ("A11", "Comparator Identity", "Deterministic vocabulary + LLM residue", "a11.comparator_identity",
     "INN/ATC looked up from the seed vocabulary first; the LLM is only consulted for names the vocabulary can't resolve."),
    ("A14", "Harmonisation (comparator grouping)", "Deterministic", "-",
     "Groups extracted comparator mentions into candidate identities. No LLM call at all despite the label -- pure dedup/grouping logic."),
    ("A12", "Scope Adjudication", "LLM prompt", "a12.scope_adjudication",
     "In/out/uncertain verdict per candidate comparator, with evidence FOR and AGAINST retained either way."),
    ("A13", "Per-State Assignment", "Deterministic", "-",
     "Derives all-27-state verdicts from which states' searches actually surfaced each comparator. No LLM."),
    ("A15", "Outcome Harmonisation + Catalog Coverage", "LLM prompts", "a14.outcome_harmonization, a16.outcome_rationale",
     "Groups outcome mentions against the fixed 11-item catalog and writes each outcome's rationale. (Runs under the orchestrator's 'A15' stage despite one prompt being named a14.*.)"),
    ("A16", "Consolidation (comparators)", "Deterministic assembly + LLM prompts", "a16.comparator_rationale, a16.indication_synthesis",
     "Assembles the final ConsolidatedComparator objects; LLM only writes the human-readable rationale/indication text, never the structured fields."),
    ("A17", "Completeness Audit", "Deterministic", "-",
     "Structural assertions only: 27-state coverage, evidence-required invariant, no-PICO-sets-in-Phase-1, etc. Fails the run rather than silently passing."),
]

HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(color="FFFFFF", bold=True)


def _write_table(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    for i, h in enumerate(headers, start=1):
        col = get_column_letter(i)
        width = max(12, min(60, len(str(h)) + 4))
        ws.column_dimensions[col].width = width


def stage_for_prompt(prompt_id: str) -> str:
    return PROMPT_TO_STAGE.get(prompt_id, "?")


def main() -> None:
    runs = {drug: json.load(open(os.path.join(OUT_DIR, f"{drug}.json"), encoding="utf-8"))
           for drug in DRUGS}

    wb = Workbook()

    # ---- Run Summary --------------------------------------------------
    ws = wb.active
    ws.title = "Run Summary"
    rows = []
    for drug, d in runs.items():
        out = d["output"]
        v = out["validation"]
        rows.append([
            drug, d["status"], d["total_latency_s"], d["total_cost_usd"],
            d["bedrock"]["total_cost_usd"], d["tavily"]["total_cost_usd"],
            len(out["comparators"]), len(out["outcomes"].get("catalog_coverage", [])),
            v["records_extracted"], v["grounded"], v["claim_validated"], v["claim_rejected"],
            v["subject_drug_rejections"], v["role_rejections"],
            v["scope_in"], v["scope_out"], v["scope_uncertain"],
            d["bedrock"]["calls"], d["bedrock"]["failed_calls"],
            d["tavily"]["calls"], d["tavily"]["failed_calls"],
        ])
    _write_table(ws, [
        "drug", "status", "total_latency_s", "total_cost_usd", "bedrock_cost_usd",
        "tavily_cost_usd", "num_comparators", "num_catalog_outcomes",
        "records_extracted", "grounded", "claim_validated", "claim_rejected",
        "subject_drug_rejections", "role_rejections", "scope_in", "scope_out",
        "scope_uncertain", "bedrock_calls", "bedrock_failed", "tavily_calls", "tavily_failed",
    ], rows)

    # ---- Comparators ----------------------------------------------------
    ws = wb.create_sheet("Comparators")
    rows = []
    for drug, d in runs.items():
        for c in d["output"]["comparators"]:
            rows.append([
                drug, c["generic_name"], c["comparator_id"], c["inn"], c["atc_code"],
                c["class_or_mechanism"], c["class_source"], c["is_combination"],
                "; ".join(c.get("components", [])), c["role"], c["indication"],
                c["line_of_therapy"], c["comparator_scenario"], c["retain_all_status"],
                c["recommendation_strength"], c["member_state_count"],
                "; ".join(c.get("member_states", [])), "; ".join(str(t) for t in c.get("tiers", [])),
                c["cross_tier_confirmed"], c["or_alternative"], c["general_evidence_flag"],
                c["rationale"], c["origin"], len(c.get("sources", [])), len(c.get("evidence", [])),
            ])
    _write_table(ws, [
        "drug", "generic_name", "comparator_id", "inn", "atc_code", "class_or_mechanism",
        "class_source", "is_combination", "components", "role", "indication",
        "line_of_therapy", "comparator_scenario", "retain_all_status", "recommendation_strength",
        "member_state_count", "member_states", "tiers", "cross_tier_confirmed", "or_alternative",
        "general_evidence_flag", "rationale", "origin", "num_sources", "num_evidence",
    ], rows)

    # ---- Per-State Verdicts ---------------------------------------------
    ws = wb.create_sheet("Per-State Verdicts")
    rows = []
    for drug, d in runs.items():
        for c in d["output"]["comparators"]:
            for v in c.get("per_member_state", []):
                rows.append([
                    drug, c["generic_name"], v["member_state"], v["verdict"], v.get("tier"),
                    v.get("source_id", ""), v.get("source_url", ""),
                    (v.get("evidence_quote") or "")[:300], v.get("reason", ""),
                ])
    _write_table(ws, [
        "drug", "generic_name", "member_state", "verdict", "tier", "source_id",
        "source_url", "evidence_quote", "reason",
    ], rows)

    # ---- Outcomes ---------------------------------------------------------
    ws = wb.create_sheet("Outcomes")
    rows = []
    for drug, d in runs.items():
        oc = d["output"]["outcomes"]
        for cat in ("clinical_effectiveness", "safety", "quality_of_life", "clinician_patient_reported"):
            for o in oc.get(cat, []):
                rows.append([
                    drug, cat, o["concept"], o["outcome_id"], o["category"], o["catalog_id"],
                    o["listed"], o["instrument"], o["unit_of_measurement"], o["requirement_type"],
                    o["coverage_status"], o["rationale"], o["origin"], len(o.get("sources", [])),
                ])
    _write_table(ws, [
        "drug", "bucket", "concept", "outcome_id", "category", "catalog_id", "listed",
        "instrument", "unit_of_measurement", "requirement_type", "coverage_status",
        "rationale", "origin", "num_sources",
    ], rows)

    # ---- Catalog Coverage ---------------------------------------------
    ws = wb.create_sheet("Catalog Coverage")
    rows = []
    for drug, d in runs.items():
        for cc in d["output"]["outcomes"].get("catalog_coverage", []):
            rows.append([drug, cc["catalog_id"], cc["display_name"], cc["category"],
                        cc["status"], cc.get("matched_outcome_id", ""), cc.get("detail", "")])
    _write_table(ws, ["drug", "catalog_id", "display_name", "category", "status",
                     "matched_outcome_id", "detail"], rows)

    # ---- By Member State --------------------------------------------------
    ws = wb.create_sheet("By Member State")
    rows = []
    for drug, d in runs.items():
        for e in d["output"]["by_member_state"]:
            rows.append([drug, e["member_state"], e["status"],
                        "; ".join(e.get("comparators", [])), e.get("state_search_status", "")])
    _write_table(ws, ["drug", "member_state", "status", "comparators", "state_search_status"], rows)

    # ---- Bedrock Calls ------------------------------------------------
    ws = wb.create_sheet("Bedrock Calls")
    rows = []
    for drug, d in runs.items():
        for c in d["bedrock"]["calls_detail"]:
            rows.append([
                drug, stage_for_prompt(c["prompt_id"]), c["prompt_id"], c["prompt_version"],
                c["input_tokens"], c["output_tokens"], c["cost_usd"], c["latency_s"],
                c["attempts"], c["error"],
            ])
    _write_table(ws, ["drug", "stage", "prompt_id", "prompt_version", "input_tokens",
                     "output_tokens", "cost_usd", "latency_s", "attempts", "error"], rows)

    # ---- Tavily Calls -------------------------------------------------
    ws = wb.create_sheet("Tavily Calls")
    rows = []
    for drug, d in runs.items():
        for c in d["tavily"]["calls_detail"]:
            rows.append([drug, c["op"], c["detail"], c["credits"], c["cost_usd"],
                        c["latency_s"], c["attempts"], c["results"], c["error"]])
    _write_table(ws, ["drug", "op", "detail", "credits", "cost_usd", "latency_s",
                     "attempts", "results", "error"], rows)

    # ---- Stage Latency --------------------------------------------------
    ws = wb.create_sheet("Stage Latency")
    rows = []
    for drug, d in runs.items():
        for stage, secs in d["stage_latency_s"].items():
            rows.append([drug, stage, secs])
    _write_table(ws, ["drug", "stage", "latency_s"], rows)

    # ---- Pipeline Steps (reference) -------------------------------------
    ws = wb.create_sheet("Pipeline Steps")
    _write_table(ws, ["stage", "name", "type", "prompt_id(s)", "description"], PIPELINE_STEPS)
    for row in ws.iter_rows(min_row=2, max_col=5):
        row[4].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["E"].width = 100

    # ---- Run Manifest ---------------------------------------------------
    ws = wb.create_sheet("Run Manifest")
    rows = []
    for drug, d in runs.items():
        m = d["output"]["run_manifest"]
        rows.append([
            drug, m.get("model"), m.get("code_version"), m.get("outcome_catalog_version"),
            json.dumps(m.get("prompt_versions", {})), m.get("tier3_enabled"),
            json.dumps(m.get("providers", {})), m.get("started_at"), m.get("finished_at"),
            m.get("duration_s"),
        ])
    _write_table(ws, ["drug", "model", "code_version", "outcome_catalog_version",
                     "prompt_versions", "tier3_enabled", "providers", "started_at",
                     "finished_at", "duration_s"], rows)

    out_path = os.path.join(OUT_DIR, "GT_Phase1_Results.xlsx")
    wb.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    main()
