"""
Tests for streamlit_app/excel_export.py covering the SME's latest ask:
print the PubMed/tier-1/2/3 query strings and the evidence_status /
data_maturity / trial_status / peer_review_status fields in the Excel
export -- both in the final-output "Evidence" sheets (already covered
before this change) and in the shared per-stage debug-trail sheets that
use _record_row()/_RECORD_HEADERS (A8/A9/A10/A10-Usable/A14/A15-Groups),
which previously dropped those four fields.
"""
import io
import sys
import unittest

sys.path.insert(0, "streamlit_app")

import openpyxl

from jca_phase1 import config as C
from streamlit_app import excel_export


def _load(run):
    wb = openpyxl.load_workbook(io.BytesIO(excel_export.build_excel(run)))
    return wb


def _headers(ws):
    return [c.value for c in ws[1]]


class TestRecordRowEvidenceFields(unittest.TestCase):
    """Looked up by header name, not trailing position -- _RECORD_HEADERS has
    grown twice since these fields were added (thin_mention, then the
    population_context/recommendation_strength/etc. batch), so pinning a
    fixed tail slice breaks on every future addition for no real reason."""

    def _row_by_header(self, row):
        return dict(zip(excel_export._RECORD_HEADERS, row))

    def test_record_row_includes_the_four_sme_requested_fields(self):
        row = excel_export._record_row({
            "finding_id": "f1", "finding_type": "comparator", "subject_drug": "DrugX",
            "member_state": "DE", "source_class": C.SRC_CONFERENCE, "tier": 2,
            "evidence_status": "conference_abstract", "data_maturity": "interim",
            "trial_status": "TERMINATED", "peer_review_status": "preprint",
        })
        by_header = self._row_by_header(row)
        self.assertEqual(by_header["evidence_status"], "conference_abstract")
        self.assertEqual(by_header["data_maturity"], "interim")
        self.assertEqual(by_header["trial_status"], "TERMINATED")
        self.assertEqual(by_header["peer_review_status"], "preprint")
        self.assertEqual(len(row), len(excel_export._RECORD_HEADERS))

    def test_record_headers_name_the_four_fields(self):
        for name in ("evidence_status", "data_maturity", "trial_status",
                    "peer_review_status"):
            self.assertIn(name, excel_export._RECORD_HEADERS)
        self.assertIn("thin_mention", excel_export._RECORD_HEADERS)


class TestQueryPlanTierColumn(unittest.TestCase):
    def _run_with_plan_and_attempts(self):
        return {
            "output": {"comparators": [], "outcomes": {}},
            "debug_capture": {
                "A6_query_plan": {"plan": [
                    {"query": "drug AND disease", "source_class": C.SRC_HTA_REGULATORY,
                     "member_state": "DE", "pass_type": "primary", "domains": [],
                     "max_urls": 5, "language": "en", "note": ""},
                    {"query": "guideline query", "source_class": C.SRC_CLINICAL_GUIDELINE,
                     "member_state": "EU-wide", "pass_type": "primary", "domains": [],
                     "max_urls": 5, "language": "en", "note": ""},
                ]},
                "A7_retrieval": {"attempts": [
                    {"member_state": "EU-wide", "source_class": C.SRC_PUBMED,
                     "attempted": True, "documents_retrieved": 10, "status": "ok",
                     "detail": "[1 drug-anchored] q1 | [2 guideline-restricted] q2 | "
                              "[3 landscape, unrestricted] q3"},
                ], "documents": []},
            },
        }

    def test_a6_query_plan_sheet_has_a_tier_column_matching_source_class(self):
        wb = _load(self._run_with_plan_and_attempts())
        ws = wb["A6 Query Plan (Mix)"]
        headers = _headers(ws)
        self.assertIn("tier", headers)
        tier_idx = headers.index("tier") + 1
        source_idx = headers.index("source_class") + 1
        for row in ws.iter_rows(min_row=2, values_only=False):
            source_class = row[source_idx - 1].value
            tier = row[tier_idx - 1].value
            self.assertEqual(tier, C.TIER_OF_SOURCE_CLASS[source_class])

    def test_a7_attempts_sheet_carries_pubmed_query_strings_and_tier(self):
        wb = _load(self._run_with_plan_and_attempts())
        ws = wb["A7 Retrieval Attempts (Det)"]
        headers = _headers(ws)
        self.assertIn("query_sent", headers)
        self.assertIn("tier", headers)
        row = list(ws.iter_rows(min_row=2, values_only=True))[0]
        detail = row[headers.index("query_sent")]
        self.assertIn("drug-anchored", detail)
        self.assertIn("guideline-restricted", detail)
        self.assertIn("landscape", detail)
        self.assertEqual(row[headers.index("tier")], C.TIER_OF_SOURCE_CLASS[C.SRC_PUBMED])


class TestHighPriorityCompletenessGapsClosed(unittest.TestCase):
    """2026-09-24 completeness audit: fields that exist in the data model but
    were printed on ZERO sheets. Covers the high-priority subset only --
    population_context/recommendation_strength/retrieval_method/doc metadata
    on the shared debug-record sheets, the AUDITED completeness matrix,
    Phase1Output.notes, Intervention's resolved identity, and
    member_state_summary."""

    def test_record_row_carries_population_context_and_doc_metadata(self):
        row = excel_export._record_row({
            "finding_id": "cmp-1", "finding_type": "comparator", "subject_drug": "DrugX",
            "recommendation_strength": "preferred", "retrieval_method": "tavily_search",
            "document_title": "ESMO Guideline", "document_date": "2024",
            "population_context": {"disease": "SCLC", "stage": "extensive-stage",
                                   "line_of_therapy": "second-line"},
        })
        by_header = dict(zip(excel_export._RECORD_HEADERS, row))
        self.assertIn("SCLC", by_header["population_context"])
        self.assertIn("extensive-stage", by_header["population_context"])
        self.assertIn("second-line", by_header["population_context"])
        self.assertEqual(by_header["recommendation_strength"], "preferred")
        self.assertEqual(by_header["retrieval_method"], "tavily_search")
        self.assertEqual(by_header["document_title"], "ESMO Guideline")
        self.assertEqual(by_header["document_date"], "2024")

    def test_record_row_carries_outcome_specific_fields_only_for_outcomes(self):
        outcome_row = excel_export._record_row({
            "finding_id": "out-1", "finding_type": "outcome",
            "outcome": {"measure": "ORR", "unit": "%", "instrument": "RECIST",
                       "is_requirement": True},
        })
        by_header = dict(zip(excel_export._RECORD_HEADERS, outcome_row))
        self.assertEqual(by_header["outcome_unit"], "%")
        self.assertEqual(by_header["outcome_instrument"], "RECIST")
        self.assertTrue(by_header["outcome_is_requirement"])

        comparator_row = excel_export._record_row({
            "finding_id": "cmp-1", "finding_type": "comparator",
            "comparator": {"as_stated": "Topotecan"},
        })
        by_header2 = dict(zip(excel_export._RECORD_HEADERS, comparator_row))
        self.assertIsNone(by_header2["outcome_unit"])

    def test_a17_source_class_matrix_sheet_reads_the_audited_completeness_field(self):
        run = {"output": {"comparators": [], "outcomes": {},
                          "completeness": {"source_class_matrix": [
                              {"member_state": "Austria", "source_class": C.SRC_HTA_REGULATORY,
                               "attempted": True, "documents_retrieved": 2,
                               "status": "found", "detail": ""}]}},
              "debug_capture": {}}
        wb = _load(run)
        self.assertIn("A17 Source Class Matrix (Det)", wb.sheetnames)
        ws = wb["A17 Source Class Matrix (Det)"]
        row = list(ws.iter_rows(min_row=2, values_only=True))[0]
        headers = _headers(ws)
        self.assertEqual(row[headers.index("member_state")], "Austria")
        self.assertEqual(row[headers.index("documents_retrieved")], 2)

    def test_run_notes_sheet_surfaces_operational_warnings(self):
        run = {"output": {"comparators": [], "outcomes": {},
                          "notes": ["[A10] Claim validation skipped: it requires both an "
                                   "LLM and a search provider."]},
              "debug_capture": {}}
        wb = _load(run)
        self.assertIn("Run Notes", wb.sheetnames)
        ws = wb["Run Notes"]
        row = list(ws.iter_rows(min_row=2, values_only=True))[0]
        self.assertIn("Claim validation skipped", row[0])

    def test_run_summary_includes_member_state_summary(self):
        run = {"output": {"comparators": [], "outcomes": {},
                          "member_state_summary": {"identified_count": 5,
                                                   "not_identified_count": 22, "total": 27}},
              "debug_capture": {}}
        wb = _load(run)
        ws = wb["Run Summary"]
        rows = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["member_states_identified"], 5)
        self.assertEqual(rows["member_states_not_identified"], 22)
        self.assertEqual(rows["member_states_total"], 27)

    def test_a2_intervention_sheet_includes_resolved_identity(self):
        run = {"output": {"comparators": [], "outcomes": {},
                          "intervention": {"fields": {}, "inn_resolved": "tovorafenib",
                                          "atc_code": "L01EX23", "user_confirmed": True}},
              "debug_capture": {}}
        wb = _load(run)
        ws = wb["A2 Intervention Fields (Mix)"]
        rows = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(rows["inn_resolved"], "tovorafenib")
        self.assertEqual(rows["atc_code"], "L01EX23")
        self.assertTrue(rows["user_confirmed"])


if __name__ == "__main__":
    unittest.main()
