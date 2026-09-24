import openpyxl
wb = openpyxl.load_workbook('fixtures/jca gt.xlsx', data_only=True)
ws = wb['Tovorafenib']
out = []
out.append(f"dims: {ws.dimensions}")
for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
    vals = [c.value for c in row]
    if any(v not in (None, '') for v in vals):
        out.append(f"{row[0].row}: {vals}")
with open('_gt_dump.txt', 'w', encoding='utf-8') as f:
    f.write('\n'.join(out))
print(len(out), 'lines')
